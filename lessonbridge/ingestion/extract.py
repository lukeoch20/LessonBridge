"""Document extraction: PDFs, Word files, spreadsheets, HTML, Markdown and text."""
from __future__ import annotations

import csv
import html
import io
import re
from pathlib import Path

CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".html": "text/html",
    ".htm": "text/html",
    ".json": "application/json",
}


def guess_content_type(path: str | Path) -> str:
    return CONTENT_TYPES.get(Path(path).suffix.lower(), "application/octet-stream")


def extract_text(data: bytes, content_type: str, name: str = "") -> str:
    ct = (content_type or "").lower()
    if not ct or ct == "application/octet-stream":
        ct = guess_content_type(name)
    try:
        if "pdf" in ct:
            return _pdf(data)
        if "wordprocessingml" in ct or name.lower().endswith(".docx"):
            return _docx(data)
        if "spreadsheetml" in ct or name.lower().endswith(".xlsx"):
            return _xlsx(data)
        if "csv" in ct:
            return _csv(data)
        if "html" in ct:
            return _html(data)
    except Exception as exc:  # pragma: no cover - defensive: never lose the upload
        return f"[extraction failed: {exc}]"
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        pages.append(f"[page {i + 1}]\n" + (page.extract_text() or ""))
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    """Paragraphs and tables in document order (LB-12); one line per table row, cells joined by ' | '.

    Heading paragraphs are prefixed with '#' so the curriculum parser can tell
    them from body text.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    out: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                text = "#" * min(int(level), 6) + " " + text
            out.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells, seen = [], set()
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat the same element
                        continue
                    seen.add(id(cell._tc))
                    cells.append(" ".join(cell.text.split()))
                if any(cells):
                    out.append(" | ".join(cells))
    return "\n".join(out)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))


def _html(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    text = re.sub(r"(?is)<(script|style|noscript).*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</h[1-6]>|</tr>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()
