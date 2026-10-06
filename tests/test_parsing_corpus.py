"""A small corpus of realistic syllabi and pacing guides (audit test recommendation 3; LB-11..14, LB-26..30)."""
from __future__ import annotations

import io

import pytest

from lessonbridge.curriculum.interpret import heuristic_extract, merge_curriculum, split_by_subject
from lessonbridge.ingestion.extract import extract_text

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def units(text, subject="english", doc_type="syllabus"):
    return [(u.title, u.quarter, u.planned_days) for u in heuristic_extract(text, subject, doc_type=doc_type).units]


def word_with_quarter_tables() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("English 7 Pacing Overview")
    for q, rows in [(1, [("Unit 1: Launching Readers and Writers", "6"), ("Unit 2: Short Fiction", "18")]), (2, [("Unit 3: Argumentative Writing", "16")]),
                    (3, [("Unit 4: Novel Study", "22")]), (4, [("Unit 5: Poetry and Media Literacy", "12")])]:
        d.add_heading(f"Quarter {q}", level=2)
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Unit", "Days"
        for a, b in rows:
            r = t.add_row()
            r.cells[0].text, r.cells[1].text = a, b
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_lb12_word_tables_keep_document_order():
    text = extract_text(word_with_quarter_tables(), DOCX, "pacing.docx")
    assert units(text) == [("Launching Readers and Writers", 1, 6), ("Short Fiction", 1, 18), ("Argumentative Writing", 2, 16), ("Novel Study", 3, 22), ("Poetry and Media Literacy", 4, 12)]


def test_lb12_xlsx_and_csv_pacing_guides():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [("Quarter", "Unit", "Title", "Weeks"), (1, 1, "Launching Readers and Writers", 2), (2, 3, "Argumentative Writing", 3)]:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    assert units(extract_text(buf.getvalue(), XLSX, "pace.xlsx")) == [("Launching Readers and Writers", 1, 10), ("Argumentative Writing", 2, 15)]
    csv = b"Quarter,Unit,Title,Days\n1,1,Launching Readers and Writers,6\n2,3,Argumentative Writing,16\n"
    assert units(extract_text(csv, "text/csv", "pace.csv")) == [("Launching Readers and Writers", 1, 6), ("Argumentative Writing", 2, 16)]


def test_lb14_dated_quarter_headings_set_the_quarter_only():
    text = "Quarter 1 (August 20 - October 30)\nUnit 1: Short Fiction (18 days)\nQuarter 2 begins November 4\nUnit 2: Argumentative Writing (16 days)\nQuarter 3 (Weeks 19-27)\nUnit 3: Novel Study\n"
    assert units(text) == [("Short Fiction", 1, 18), ("Argumentative Writing", 2, 16), ("Novel Study", 3, None)]


def test_lb26_week_by_week_plans_merge_weeks_and_never_invent_assessments():
    text = "Quarter 1\nWeeks 1-2 (Aug 20 - Aug 31): Launching Readers and Writers\nWeek 3: Short Fiction\nWeek 4: Short Fiction\nWeek 5 - Theme\nWeeks 6-9: Narrative Writing\n"
    ctx = heuristic_extract(text, "english")
    assert [(u.title, u.planned_days) for u in ctx.units] == [("Launching Readers and Writers", 10), ("Short Fiction", 10), ("Theme", 5), ("Narrative Writing", 20)]
    merged = merge_curriculum("english", ctx).spec
    theme = next(u for u in merged.units if u.title == "Theme")
    assert theme.lessons == []  # no borrowed Novel Study unit, no invented assessment


def test_lb27_a_stated_ten_day_length_is_kept():
    ctx = heuristic_extract("Unit 1: Novel Study (10 days)\nUnit 2: Research and Source Evaluation (20 days)\n", "english")
    spec = merge_curriculum("english", ctx).spec
    lengths = {u.title: u.planned_days for u in spec.units}
    assert lengths["Novel Study"] == 10 and lengths["Research and Source Evaluation"] == 20


@pytest.mark.parametrize("text", [
    "## Quarter 1\n### Unit 1: Launching Readers and Writers (6 days)\n1. Unit 2: Short Fiction (18 days)\n## Quarter 2\n- **Unit 3: Argumentative Writing** — 3 weeks\n",
    "| Quarter | Unit | Days |\n|---|---|---|\n| 1 | Launching Readers and Writers | 6 |\n| 1 | Short Fiction | 18 |\n| 2 | Argumentative Writing | 15 |\n",
])
def test_lb28_markdown_headings_lists_and_tables(text):
    assert units(text) == [("Launching Readers and Writers", 1, 6), ("Short Fiction", 1, 18), ("Argumentative Writing", 2, 15)]


def test_lb11_civics_units_are_matched_by_whole_title():
    text = open("examples/teacher_docs/civics_syllabus.md").read()
    spec = merge_curriculum("civics", heuristic_extract(text, "civics")).spec
    by_title = {u.title: u for u in spec.units}
    econ = [l.slug for l in by_title["Economics"].lessons]
    review = [l.slug for l in by_title["Civics & Economics SOL Review"].lessons]
    assert "econ-scarcity" in econ and "econ-test" in econ
    assert review and all(s.startswith("sol-") for s in review)
    assert not any(u.title.startswith("Economics:") for u in spec.units)  # the default economics unit is not duplicated at year end
    assert not any(l.slug.endswith("-assessment") for u in spec.units for l in u.lessons)  # nothing invented


def test_lb13_omitted_default_units_stay_in_their_quarter():
    text = open("examples/teacher_docs/english7_syllabus.md").read().replace("Unit 1: Launching Readers and Writers (6 days)\n", "")
    result = merge_curriculum("english", heuristic_extract(text, "english"))
    order = [(u.title, u.quarter) for u in result.spec.units]
    assert order[0][0].startswith("Launching") and order[0][1] == 1
    assert [q for _, q in order] == sorted(q for _, q in order)
    assert any("not in your syllabus" in n for n in result.notes)


def test_lb29_combined_documents_are_split_by_subject():
    text = "# English 7\nQuarter 1\nUnit 1: Short Fiction (18 days)\n# Civics and Economics\nQuarter 1\nUnit 1: Founding Documents (14 days)\n"
    parts = split_by_subject(text)
    assert units(parts["english"]) == [("Short Fiction", 1, 18)]
    assert units(parts["civics"], "civics") == [("Founding Documents", 1, 14)]


def test_lb30_imperative_procedures_become_rules():
    text = "# Substitute procedures\nTake attendance on the clipboard by the door.\nDo not allow students to leave without a pass.\nCollect cell phones in the caddy at the start of class.\n"
    rules = [r.text for r in heuristic_extract(text, "english", doc_type="classroom_procedures").rules]
    assert len(rules) == 3 and not any("procedures" in r.lower() and len(r.split()) < 4 for r in rules)


def test_grade_weighting_is_not_a_rule_but_late_work_is():
    rules = [r.text for r in heuristic_extract(open("examples/teacher_docs/english7_syllabus.md").read(), "english").rules]
    assert any(r.startswith("Late work") for r in rules)
    assert not any("weighted" in r for r in rules)
